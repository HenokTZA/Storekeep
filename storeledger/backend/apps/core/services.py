import uuid
from datetime import datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from django.conf import settings
from django.core.files.base import ContentFile
from django.db import IntegrityError, transaction
from django.db.models import Count, F, Sum
from django.db.models.functions import Coalesce
from django.db.models.functions import TruncDate
from django.utils import timezone
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont, PngImagePlugin
from reportlab.lib import colors as pdf_colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from rest_framework.exceptions import ValidationError

from .models import (
    AuditEvent,
    Expense,
    ExpenseCategory,
    FinancialTransaction,
    InventoryBalance,
    Notification,
    Party,
    Payment,
    Product,
    Purchase,
    PurchaseItem,
    Report,
    SMSLog,
    Sale,
    SaleItem,
    StockMovement,
    StoreSettings,
)


ZERO = Decimal("0.00")
AGENT_DISCOUNT_RATE = Decimal("0.015")


def money(value) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def agent_selling_price(value) -> Decimal:
    """Return the default Agent price: 1.5% below the Product selling price."""
    return money(Decimal(str(value)) * (Decimal("1.00") - AGENT_DISCOUNT_RATE))


def quantity(value) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)


class NotificationService:
    @staticmethod
    def sync_low_stock(product: Product, current_quantity: Decimal):
        event_key = f"low-stock:{product.id}"
        if current_quantity <= product.minimum_stock_threshold:
            Notification.objects.update_or_create(
                store=product.store,
                event_key=event_key,
                defaults={
                    "notification_type": Notification.NotificationType.LOW_STOCK,
                    "title": "Low Stock",
                    "message": f"{product.name} — {current_quantity} {product.unit} remaining",
                    "entity_type": "product",
                    "entity_id": str(product.id),
                    "is_read": False,
                    "is_resolved": False,
                },
            )
        else:
            Notification.objects.filter(store=product.store, event_key=event_key).update(is_resolved=True)


class InventoryService:
    @staticmethod
    @transaction.atomic
    def adjust(*, store, product, delta, movement_type, user, note="", reference_type="", reference_id=""):
        if product.store_id != store.id:
            raise ValidationError("Product does not belong to this store.")
        delta = quantity(delta)
        balance, _ = InventoryBalance.objects.select_for_update().get_or_create(product=product)
        new_quantity = balance.quantity + delta
        store_settings, _ = StoreSettings.objects.get_or_create(store=store)
        if store_settings.prevent_negative_inventory and new_quantity < 0:
            raise ValidationError({"quantity": f"Insufficient stock. Available: {balance.quantity}."})
        balance.quantity = new_quantity
        balance.save(update_fields=["quantity", "updated_at"])
        movement = StockMovement.objects.create(
            store=store,
            product=product,
            factory_name=product.factory.name,
            movement_type=movement_type,
            quantity_delta=delta,
            balance_after=new_quantity,
            reference_type=reference_type,
            reference_id=str(reference_id) if reference_id else "",
            note=note,
            created_by=user,
        )
        NotificationService.sync_low_stock(product, new_quantity)
        return movement


class LedgerService:
    @staticmethod
    def append(*, store, party, transaction_type, delta, description, user, sale=None, payment=None, purchase=None, note=""):
        delta = money(delta)
        party.current_balance = money(party.current_balance + delta)
        party.save(update_fields=["current_balance", "updated_at"])
        return FinancialTransaction.objects.create(
            store=store,
            party=party,
            transaction_type=transaction_type,
            description=description,
            delta=delta,
            running_balance=party.current_balance,
            sale=sale,
            payment=payment,
            purchase=purchase,
            note=note,
            created_by=user,
        )


class SaleService:
    @staticmethod
    @transaction.atomic
    def create(*, store, user, customer_id, amount_paid, items, idempotency_key, note=""):
        existing = Sale.objects.filter(store=store, idempotency_key=idempotency_key).prefetch_related("items").first()
        if existing:
            return existing, False
        if not items:
            raise ValidationError({"items": "At least one product is required."})

        customer = None
        if customer_id:
            customer = Party.objects.select_for_update().filter(store=store, is_active=True, id=customer_id).first()
            if not customer:
                raise ValidationError({"customer_id": "Customer was not found in this store."})
            if customer.party_type == Party.PartyType.FACTORY:
                raise ValidationError({"customer_id": "Factories are purchase suppliers and cannot be sale customers."})

        seen_products = set()
        prepared_items = []
        total = ZERO
        for item in items:
            product_id = item.get("product_id")
            if product_id in seen_products:
                raise ValidationError({"items": "A product may only appear once in a sale."})
            seen_products.add(product_id)
            product = Product.objects.select_related("factory").filter(store=store, is_active=True, id=product_id).first()
            if not product:
                raise ValidationError({"items": f"Product {product_id} was not found."})
            item_quantity = quantity(item.get("quantity", 0))
            if item_quantity <= 0:
                raise ValidationError({"items": "Every quantity must be greater than zero."})
            if item.get("unit_price") is not None:
                # A negotiated price is a snapshot for this sale item only. It
                # applies to Traders, Agents and walk-in customers, and never
                # changes the product's standard selling price.
                unit_price = money(item["unit_price"])
            elif customer and customer.party_type == Party.PartyType.AGENT:
                unit_price = agent_selling_price(product.selling_price)
            else:
                # Walk-in and Trader sales default to the standard price but
                # may explicitly override it above for one transaction.
                unit_price = money(product.selling_price)
            if unit_price < 0:
                raise ValidationError({"items": "Unit price cannot be negative."})
            pieces_per_unit = product.pieces_per_unit
            line_total = money(item_quantity * pieces_per_unit * unit_price)
            unit_cost = money(product.purchase_price)
            line_cost = money(item_quantity * pieces_per_unit * unit_cost)
            gross_profit = money(line_total - line_cost)
            total += line_total
            prepared_items.append(
                (product, item_quantity, pieces_per_unit, unit_price, line_total, unit_cost, line_cost, gross_profit)
            )

        total = money(total)
        amount_paid = money(amount_paid)
        if amount_paid < 0:
            raise ValidationError({"amount_paid": "Amount paid cannot be negative."})
        if customer is None and amount_paid != total:
            raise ValidationError({"amount_paid": "Walk-in sales must be paid in full."})

        sale = Sale.objects.create(
            store=store,
            customer=customer,
            idempotency_key=idempotency_key,
            subtotal=total,
            total=total,
            amount_paid=amount_paid,
            outstanding=money(total - amount_paid),
            note=note,
            created_by=user,
        )

        for product, item_quantity, pieces_per_unit, unit_price, line_total, unit_cost, line_cost, gross_profit in prepared_items:
            SaleItem.objects.create(
                sale=sale,
                product=product,
                product_name=product.name,
                factory_name=product.factory.name,
                quantity=item_quantity,
                pieces_per_unit=pieces_per_unit,
                unit_price=unit_price,
                line_total=line_total,
                unit_cost=unit_cost,
                line_cost=line_cost,
                gross_profit=gross_profit,
            )
            InventoryService.adjust(
                store=store,
                product=product,
                delta=-item_quantity,
                movement_type=StockMovement.MovementType.SALE,
                user=user,
                note=f"Sale {sale.id}",
                reference_type="sale",
                reference_id=sale.id,
            )

        if customer:
            LedgerService.append(
                store=store,
                party=customer,
                transaction_type=FinancialTransaction.TransactionType.SALE,
                delta=total,
                description=f"Sale {str(sale.id)[:8]}",
                user=user,
                sale=sale,
                note=note,
            )
            if amount_paid > 0:
                payment = Payment.objects.create(
                    store=store,
                    party=customer,
                    amount=amount_paid,
                    direction=Payment.Direction.RECEIVED,
                    payment_date=timezone.localdate(),
                    method=Payment.Method.CASH,
                    note=f"Payment recorded with sale {sale.id}",
                    sale=sale,
                    idempotency_key=uuid.uuid5(uuid.NAMESPACE_URL, f"sale-payment:{store.id}:{sale.id}"),
                    created_by=user,
                )
                LedgerService.append(
                    store=store,
                    party=customer,
                    transaction_type=FinancialTransaction.TransactionType.PAYMENT,
                    delta=-amount_paid,
                    description=f"Payment with sale {str(sale.id)[:8]}",
                    user=user,
                    sale=sale,
                    payment=payment,
                )

        AuditEvent.objects.create(
            store=store,
            user=user,
            action="sale.created",
            entity_type="sale",
            entity_id=str(sale.id),
            details={"total": str(total), "amount_paid": str(amount_paid)},
        )
        return sale, True


class SaleInvoiceService:
    """Render an immutable sale snapshot as a customer-facing PDF invoice."""

    @staticmethod
    def invoice_number(sale: Sale) -> str:
        local_created = timezone.localtime(sale.created_at, ZoneInfo(sale.store.timezone))
        return f"SL-{local_created:%Y%m%d}-{str(sale.id)[:8].upper()}"

    @staticmethod
    def filename(sale: Sale) -> str:
        return f"StoreLedger-Invoice-{SaleInvoiceService.invoice_number(sale)}.pdf"

    @staticmethod
    def render_pdf(sale: Sale) -> bytes:
        sale = (
            Sale.objects.select_related("store", "customer", "created_by")
            .prefetch_related("items")
            .get(pk=sale.pk)
        )
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=A4)
        page_width, page_height = A4
        left = 42
        right = page_width - 42
        invoice_number = SaleInvoiceService.invoice_number(sale)
        local_created = timezone.localtime(sale.created_at, ZoneInfo(sale.store.timezone))
        page_number = 0

        def clean(value, limit=48):
            text = str(value or "").replace("\n", " ").strip()
            return text if len(text) <= limit else f"{text[: limit - 1]}…"

        def amount(value):
            return f"{money(value):,.2f} {sale.store.currency}"

        def page_header():
            nonlocal page_number
            page_number += 1
            pdf.setFillColor(pdf_colors.HexColor("#176B54"))
            pdf.rect(0, page_height - 108, page_width, 108, fill=1, stroke=0)
            pdf.setFillColor(pdf_colors.white)
            pdf.setFont("Helvetica-Bold", 22)
            pdf.drawString(left, page_height - 48, "INVOICE")
            pdf.setFont("Helvetica-Bold", 13)
            pdf.drawRightString(right, page_height - 45, clean(sale.store.name, 40))
            pdf.setFont("Helvetica", 8.5)
            if sale.store.phone:
                pdf.drawRightString(right, page_height - 61, clean(sale.store.phone, 44))
            if sale.store.address:
                pdf.drawRightString(right, page_height - 76, clean(sale.store.address, 52))
            pdf.setFillColor(pdf_colors.HexColor("#16251F"))
            pdf.setFont("Helvetica-Bold", 9)
            pdf.drawString(left, page_height - 132, f"Invoice: {invoice_number}")
            pdf.setFont("Helvetica", 9)
            pdf.drawString(left, page_height - 148, f"Date: {local_created:%Y-%m-%d %H:%M}")
            pdf.drawRightString(right, page_height - 132, f"Status: {sale.status.upper()}")
            pdf.drawRightString(right, page_height - 148, f"Cashier: {clean(sale.created_by.get_username(), 25)}")

        def table_header(y):
            pdf.setFillColor(pdf_colors.HexColor("#E7F3EE"))
            pdf.roundRect(left, y - 17, right - left, 24, 4, fill=1, stroke=0)
            pdf.setFillColor(pdf_colors.HexColor("#176B54"))
            pdf.setFont("Helvetica-Bold", 7.5)
            pdf.drawString(left + 5, y - 8, "PRODUCT")
            pdf.drawRightString(308, y - 8, "UNITS")
            pdf.drawRightString(365, y - 8, "PCS/UNIT")
            pdf.drawRightString(423, y - 8, "PIECES")
            pdf.drawRightString(490, y - 8, "PRICE/PC")
            pdf.drawRightString(right - 5, y - 8, "AMOUNT")
            return y - 28

        page_header()
        y = page_height - 178
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawString(left, y, "Bill to")
        pdf.setFont("Helvetica", 9)
        customer_name = sale.customer.name if sale.customer else "Walk-in customer"
        pdf.drawString(left, y - 16, clean(customer_name, 58))
        if sale.customer and sale.customer.phone:
            pdf.drawString(left, y - 31, clean(sale.customer.phone, 40))
        y -= 58
        y = table_header(y)

        for item in sale.items.all():
            if y < 155:
                pdf.setFont("Helvetica", 7)
                pdf.setFillColor(pdf_colors.HexColor("#63736D"))
                pdf.drawCentredString(page_width / 2, 24, f"{invoice_number} · Page {page_number}")
                pdf.showPage()
                page_header()
                y = table_header(page_height - 178)
            total_pieces = item.quantity * item.pieces_per_unit
            pdf.setFillColor(pdf_colors.HexColor("#16251F"))
            pdf.setFont("Helvetica-Bold", 8.5)
            pdf.drawString(left + 5, y, clean(item.product_name, 34))
            pdf.setFont("Helvetica", 7.5)
            pdf.drawString(left + 5, y - 12, clean(f"Factory: {item.factory_name}", 34))
            pdf.drawRightString(308, y - 2, f"{item.quantity:g}")
            pdf.drawRightString(365, y - 2, str(item.pieces_per_unit))
            pdf.drawRightString(423, y - 2, f"{total_pieces:g}")
            pdf.drawRightString(490, y - 2, f"{money(item.unit_price):,.2f}")
            pdf.setFont("Helvetica-Bold", 8)
            pdf.drawRightString(right - 5, y - 2, f"{money(item.line_total):,.2f}")
            pdf.setStrokeColor(pdf_colors.HexColor("#D7E1DD"))
            pdf.line(left, y - 20, right, y - 20)
            y -= 34

        if y < 154:
            pdf.setFont("Helvetica", 7)
            pdf.setFillColor(pdf_colors.HexColor("#63736D"))
            pdf.drawCentredString(page_width / 2, 24, f"{invoice_number} · Page {page_number}")
            pdf.showPage()
            page_header()
            pdf.setFillColor(pdf_colors.HexColor("#16251F"))
            pdf.setFont("Helvetica-Bold", 12)
            pdf.drawString(left, page_height - 178, "Invoice totals")
            y = page_height - 198

        summary_top = y - 8
        total_balance = TransactionReceiptService._balance_after(
            sale,
            sale.outstanding if sale.customer else ZERO,
        )
        pdf.setFillColor(pdf_colors.HexColor("#F4F8F6"))
        pdf.roundRect(310, summary_top - 116, right - 310, 116, 7, fill=1, stroke=0)
        pdf.setFillColor(pdf_colors.HexColor("#4F625B"))
        pdf.setFont("Helvetica", 9)
        pdf.drawString(324, summary_top - 20, "Subtotal")
        pdf.drawString(324, summary_top - 41, "Amount paid")
        pdf.drawString(324, summary_top - 62, "This sale outstanding")
        total_balance_label = "Total outstanding" if total_balance >= 0 else "Total customer credit"
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawString(324, summary_top - 88, total_balance_label)
        pdf.setFillColor(pdf_colors.HexColor("#16251F"))
        pdf.drawRightString(right - 12, summary_top - 20, amount(sale.subtotal))
        pdf.drawRightString(right - 12, summary_top - 41, amount(sale.amount_paid))
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawRightString(right - 12, summary_top - 62, amount(sale.outstanding))
        pdf.setFillColor(pdf_colors.HexColor("#A83B32") if total_balance > 0 else pdf_colors.HexColor("#176B54"))
        pdf.drawRightString(right - 12, summary_top - 88, amount(abs(total_balance)))
        pdf.setFont("Helvetica", 7.5)
        pdf.setFillColor(pdf_colors.HexColor("#63736D"))
        pdf.drawString(left, summary_top - 20, "Quantity is shown in sale units/packs.")
        pdf.drawString(left, summary_top - 34, "Piece price × pieces per unit × units = line amount.")
        if sale.note:
            pdf.drawString(left, summary_top - 55, f"Note: {clean(sale.note, 62)}")
        pdf.drawCentredString(page_width / 2, 24, f"{invoice_number} · Page {page_number}")
        pdf.save()
        return buffer.getvalue()


class TransactionReceiptService:
    """Render sale and purchase snapshots as high-resolution PNG receipts."""

    WIDTH = 1240
    PRIMARY = "#176B54"
    PRIMARY_SOFT = "#E7F3EE"
    BACKGROUND = "#F4F8F6"
    TEXT = "#16251F"
    MUTED = "#52645D"
    BORDER = "#D7E1DD"
    DANGER = "#A83B32"
    AMHARIC_LABELS = {
        "SALE RECEIPT": "የሽያጭ ደረሰኝ",
        "PURCHASE RECEIPT": "የግዢ ደረሰኝ",
        "RECEIPT NUMBER": "የደረሰኝ ቁጥር",
        "DATE": "ቀን",
        "STATUS": "ሁኔታ",
        "RECORDED BY": "የመዘገበው",
        "COMPLETED": "ተጠናቋል",
        "BILL TO": "ለ",
        "PURCHASED FROM": "የተገዛበት",
        "PRODUCT DETAILS": "የምርት ዝርዝር",
        "SALE AMOUNT": "የሽያጭ መጠን",
        "PURCHASE AMOUNT": "የግዢ መጠን",
        "Factory": "ፋብሪካ",
        "units": "ዩኒቶች",
        "pcs/unit": "ቁራጭ/ዩኒት",
        "pc": "ቁራጭ",
        "total pieces": "ጠቅላላ ቁራጮች",
        "TOTAL": "ጠቅላላ",
        "AMOUNT PAID": "የተከፈለ",
        "THIS SALE OUTSTANDING": "የዚህ ሽያጭ ቀሪ ዕዳ",
        "THIS PURCHASE OUTSTANDING": "የዚህ ግዢ ቀሪ ዕዳ",
        "TOTAL OUTSTANDING": "ጠቅላላ ቀሪ ዕዳ",
        "TOTAL CUSTOMER CREDIT": "ጠቅላላ የደንበኛ ክሬዲት",
        "TOTAL OWED TO FACTORY": "ለፋብሪካ ጠቅላላ የሚከፈል",
        "TOTAL FACTORY OWES": "ፋብሪካው ጠቅላላ የሚከፍለው",
        "REFERENCE / INVOICE": "ማጣቀሻ / ደረሰኝ",
        "NOTE": "ማስታወሻ",
        "Generated securely by StoreLedger": "በStoreLedger በደህንነት የተዘጋጀ",
        "Walk-in customer": "ያልተመዘገበ ደንበኛ",
    }

    @staticmethod
    def sale_number(sale: Sale) -> str:
        local_created = timezone.localtime(sale.created_at, ZoneInfo(sale.store.timezone))
        return f"SL-{local_created:%Y%m%d}-{str(sale.id)[:8].upper()}"

    @staticmethod
    def purchase_number(purchase: Purchase) -> str:
        return f"PL-{purchase.purchase_date:%Y%m%d}-{str(purchase.id)[:8].upper()}"

    @staticmethod
    def filename(kind: str, document) -> str:
        if kind == "sale":
            number = TransactionReceiptService.sale_number(document)
            label = "Sale"
        else:
            number = TransactionReceiptService.purchase_number(document)
            label = "Purchase"
        return f"StoreLedger-{label}-Receipt-{number}.png"

    @staticmethod
    def _font(size: int, bold: bool = False):
        filename = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
        windows_filename = "arialbd.ttf" if bold else "arial.ttf"
        candidates = (
            Path(settings.BASE_DIR) / "assets" / "fonts" / "NotoSansEthiopic.ttf",
            Path(settings.BASE_DIR) / "assets" / "fonts" / filename,
            Path("/usr/share/fonts/truetype/dejavu") / filename,
            Path("/usr/share/fonts/dejavu") / filename,
            Path("C:/Windows/Fonts") / windows_filename,
        )
        for candidate in candidates:
            try:
                font = ImageFont.truetype(str(candidate), size=size)
                if candidate.name == "NotoSansEthiopic.ttf":
                    try:
                        font.set_variation_by_name("Bold" if bold else "Regular")
                    except (AttributeError, OSError):
                        pass
                return font
            except OSError:
                continue
        try:
            return ImageFont.load_default(size=size)
        except TypeError:  # Pillow versions before the scalable default font.
            return ImageFont.load_default()

    @staticmethod
    def _label(value: str, language: str) -> str:
        if language == "am":
            return TransactionReceiptService.AMHARIC_LABELS.get(value, value)
        return value

    @staticmethod
    def _quantity_text(value) -> str:
        result = f"{Decimal(str(value)):f}".rstrip("0").rstrip(".")
        return result or "0"

    @staticmethod
    def _balance_after(document, fallback) -> Decimal:
        entry = document.ledger_entries.order_by("-created_at", "-id").first()
        return money(entry.running_balance if entry else fallback)

    @staticmethod
    def _wrap(draw, text, font, max_width):
        words = str(text or "").replace("\n", " ").split()
        if not words:
            return []
        lines = []
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
        return lines

    @staticmethod
    def _render(*, kind, store, number, counterparty_name, counterparty_phone, created_at,
                cashier, status, items, total, amount_paid, outstanding, total_balance,
                reference="", note="", language="en") -> bytes:
        width = TransactionReceiptService.WIDTH
        margin = 64
        row_height = 132
        probe = Image.new("RGB", (width, 100), "white")
        probe_draw = ImageDraw.Draw(probe)
        note_font = TransactionReceiptService._font(25, bold=True)
        note_lines = TransactionReceiptService._wrap(probe_draw, note, note_font, width - (margin * 2) - 48)[:4]
        notes_height = (len(note_lines) * 36 + 94 if note_lines else 0) + (76 if reference else 0)
        height = max(1420, 1220 + len(items) * row_height + notes_height)
        image = Image.new("RGBA", (width, height), TransactionReceiptService.BACKGROUND)
        draw = ImageDraw.Draw(image)

        title_font = TransactionReceiptService._font(50, bold=True)
        store_font = TransactionReceiptService._font(35, bold=True)
        heading_font = TransactionReceiptService._font(29, bold=True)
        label_font = TransactionReceiptService._font(21, bold=True)
        body_font = TransactionReceiptService._font(25, bold=True)
        small_font = TransactionReceiptService._font(20, bold=True)
        amount_font = TransactionReceiptService._font(30, bold=True)
        total_font = TransactionReceiptService._font(35, bold=True)

        draw.rectangle((0, 0, width, 238), fill=TransactionReceiptService.PRIMARY)
        title = TransactionReceiptService._label("SALE RECEIPT" if kind == "sale" else "PURCHASE RECEIPT", language)
        draw.text((margin, 66), title, fill="white", font=title_font)
        store_name = str(store.name)[:42]
        store_box = draw.textbbox((0, 0), store_name, font=store_font)
        if store_box[2] - store_box[0] > 620:
            store_font = TransactionReceiptService._font(28, bold=True)
            store_box = draw.textbbox((0, 0), store_name, font=store_font)
        draw.text((width - margin - (store_box[2] - store_box[0]), 72), store_name, fill="white", font=store_font)
        if store.phone:
            phone_box = draw.textbbox((0, 0), store.phone, font=small_font)
            draw.text((width - margin - (phone_box[2] - phone_box[0]), 126), store.phone, fill="white", font=small_font)
        if store.address:
            address = str(store.address).replace("\n", " ")[:64]
            address_box = draw.textbbox((0, 0), address, font=small_font)
            draw.text((width - margin - (address_box[2] - address_box[0]), 164), address, fill="white", font=small_font)

        y = 276
        draw.rounded_rectangle((margin, y, width - margin, y + 226), radius=24, fill="white", outline=TransactionReceiptService.BORDER, width=2)
        draw.text((margin + 30, y + 28), TransactionReceiptService._label("RECEIPT NUMBER", language), fill=TransactionReceiptService.MUTED, font=label_font)
        draw.text((margin + 30, y + 65), number, fill=TransactionReceiptService.TEXT, font=body_font)
        draw.text((margin + 30, y + 116), TransactionReceiptService._label("DATE", language), fill=TransactionReceiptService.MUTED, font=label_font)
        draw.text((margin + 30, y + 153), created_at, fill=TransactionReceiptService.TEXT, font=body_font)
        right_x = width // 2 + 52
        draw.text((right_x, y + 28), TransactionReceiptService._label("STATUS", language), fill=TransactionReceiptService.MUTED, font=label_font)
        draw.text((right_x, y + 65), TransactionReceiptService._label(status.upper(), language), fill=TransactionReceiptService.PRIMARY, font=body_font)
        draw.text((right_x, y + 116), TransactionReceiptService._label("RECORDED BY", language), fill=TransactionReceiptService.MUTED, font=label_font)
        draw.text((right_x, y + 153), cashier, fill=TransactionReceiptService.TEXT, font=body_font)

        y += 258
        party_label = TransactionReceiptService._label("BILL TO" if kind == "sale" else "PURCHASED FROM", language)
        draw.text((margin, y), party_label, fill=TransactionReceiptService.PRIMARY, font=label_font)
        draw.text((margin, y + 38), counterparty_name, fill=TransactionReceiptService.TEXT, font=heading_font)
        if counterparty_phone:
            draw.text((margin, y + 80), counterparty_phone, fill=TransactionReceiptService.MUTED, font=small_font)
        y += 130

        items_top = y
        items_bottom = y + 72 + len(items) * row_height
        draw.rounded_rectangle((margin, items_top, width - margin, items_bottom), radius=24, fill="white", outline=TransactionReceiptService.BORDER, width=2)
        draw.rounded_rectangle((margin, items_top, width - margin, items_top + 72), radius=24, fill=TransactionReceiptService.PRIMARY_SOFT)
        draw.rectangle((margin, items_top + 48, width - margin, items_top + 72), fill=TransactionReceiptService.PRIMARY_SOFT)
        draw.text((margin + 28, items_top + 20), TransactionReceiptService._label("PRODUCT DETAILS", language), fill=TransactionReceiptService.PRIMARY, font=label_font)
        amount_header = TransactionReceiptService._label("SALE AMOUNT" if kind == "sale" else "PURCHASE AMOUNT", language)
        amount_header_box = draw.textbbox((0, 0), amount_header, font=label_font)
        draw.text((width - margin - 28 - (amount_header_box[2] - amount_header_box[0]), items_top + 20), amount_header, fill=TransactionReceiptService.PRIMARY, font=label_font)
        row_y = items_top + 72
        for index, item in enumerate(items):
            if index:
                draw.line((margin + 24, row_y, width - margin - 24, row_y), fill=TransactionReceiptService.BORDER, width=2)
            product_name = str(item["product_name"])[:52]
            draw.text((margin + 28, row_y + 18), product_name, fill=TransactionReceiptService.TEXT, font=body_font)
            amount_text = f"{money(item['line_total']):,.2f} {store.currency}"
            amount_box = draw.textbbox((0, 0), amount_text, font=amount_font)
            draw.text((width - margin - 28 - (amount_box[2] - amount_box[0]), row_y + 17), amount_text, fill=TransactionReceiptService.TEXT, font=amount_font)
            factory_label = TransactionReceiptService._label("Factory", language)
            draw.text((margin + 28, row_y + 59), f"{factory_label}: {str(item['factory_name'])[:44]}", fill=TransactionReceiptService.PRIMARY, font=small_font)
            units = TransactionReceiptService._quantity_text(item["quantity"])
            pieces = TransactionReceiptService._quantity_text(Decimal(str(item["quantity"])) * int(item["pieces_per_unit"]))
            units_label = TransactionReceiptService._label("units", language)
            pack_label = TransactionReceiptService._label("pcs/unit", language)
            piece_label = TransactionReceiptService._label("pc", language)
            total_pieces_label = TransactionReceiptService._label("total pieces", language)
            calculation = f"{units} {units_label} x {item['pieces_per_unit']} {pack_label} x {money(item['unit_rate']):,.2f} {store.currency}/{piece_label}"
            draw.text((margin + 28, row_y + 94), f"{calculation}  |  {pieces} {total_pieces_label}", fill=TransactionReceiptService.MUTED, font=small_font)
            row_y += row_height
        y = items_bottom + 34

        summary_height = 302
        draw.rounded_rectangle((margin, y, width - margin, y + summary_height), radius=24, fill="white", outline=TransactionReceiptService.BORDER, width=2)
        transaction_label = TransactionReceiptService._label("THIS SALE OUTSTANDING" if kind == "sale" else "THIS PURCHASE OUTSTANDING", language)
        labels = [TransactionReceiptService._label("TOTAL", language), TransactionReceiptService._label("AMOUNT PAID", language), transaction_label]
        values = [total, amount_paid, abs(outstanding)]
        for index, (label, value) in enumerate(zip(labels, values)):
            row = y + 30 + index * 58
            draw.text((margin + 30, row), label, fill=TransactionReceiptService.MUTED, font=label_font)
            value_text = f"{money(value):,.2f} {store.currency}"
            value_box = draw.textbbox((0, 0), value_text, font=amount_font)
            draw.text((width - margin - 30 - (value_box[2] - value_box[0]), row - 4), value_text, fill=TransactionReceiptService.TEXT, font=amount_font)
        if kind == "sale":
            balance_label = TransactionReceiptService._label("TOTAL OUTSTANDING" if total_balance >= 0 else "TOTAL CUSTOMER CREDIT", language)
            balance_color = TransactionReceiptService.DANGER if total_balance > 0 else TransactionReceiptService.PRIMARY
        else:
            balance_label = TransactionReceiptService._label("TOTAL OWED TO FACTORY" if total_balance <= 0 else "TOTAL FACTORY OWES", language)
            balance_color = TransactionReceiptService.DANGER if total_balance < 0 else TransactionReceiptService.PRIMARY
        draw.rounded_rectangle((margin + 22, y + 204, width - margin - 22, y + 278), radius=18, fill=TransactionReceiptService.PRIMARY_SOFT)
        draw.text((margin + 46, y + 226), balance_label, fill=balance_color, font=body_font)
        balance_text = f"{abs(money(total_balance)):,.2f} {store.currency}"
        balance_box = draw.textbbox((0, 0), balance_text, font=total_font)
        draw.text((width - margin - 46 - (balance_box[2] - balance_box[0]), y + 218), balance_text, fill=balance_color, font=total_font)
        y += summary_height + 28

        if reference:
            draw.text((margin, y), TransactionReceiptService._label("REFERENCE / INVOICE", language), fill=TransactionReceiptService.PRIMARY, font=label_font)
            draw.text((margin, y + 36), str(reference)[:90], fill=TransactionReceiptService.TEXT, font=body_font)
            y += 76
        if note_lines:
            draw.text((margin, y), TransactionReceiptService._label("NOTE", language), fill=TransactionReceiptService.PRIMARY, font=label_font)
            y += 38
            for line in note_lines:
                draw.text((margin, y), line, fill=TransactionReceiptService.TEXT, font=note_font)
                y += 36

        footer_y = height - 76
        draw.line((margin, footer_y - 26, width - margin, footer_y - 26), fill=TransactionReceiptService.BORDER, width=2)
        footer = TransactionReceiptService._label("Generated securely by StoreLedger", language)
        footer_box = draw.textbbox((0, 0), footer, font=small_font)
        draw.text(((width - (footer_box[2] - footer_box[0])) / 2, footer_y), footer, fill=TransactionReceiptService.MUTED, font=small_font)

        watermark_font = TransactionReceiptService._font(82, bold=True)
        watermark_text = store.name.upper()[:42]
        watermark_box = draw.textbbox((0, 0), watermark_text, font=watermark_font)
        tile = Image.new("RGBA", (max(900, watermark_box[2] + 120), 230), (255, 255, 255, 0))
        tile_draw = ImageDraw.Draw(tile)
        tile_draw.text((60, 60), watermark_text, font=watermark_font, fill=(23, 107, 84, 18))
        rotated = tile.rotate(24, expand=True, resample=Image.Resampling.BICUBIC)
        for watermark_y in range(290, height - 150, 520):
            image.alpha_composite(rotated, ((width - rotated.width) // 2, watermark_y))

        output = BytesIO()
        metadata = PngImagePlugin.PngInfo()
        metadata.add_text("Store", store.name)
        metadata.add_text("Receipt", number)
        metadata.add_text("Transaction Type", kind)
        metadata.add_text("Transaction Outstanding", str(money(outstanding)))
        metadata.add_text("Account Balance After", str(money(total_balance)))
        metadata.add_text("Language", language)
        image.convert("RGB").save(output, format="PNG", optimize=True, pnginfo=metadata)
        return output.getvalue()

    @staticmethod
    def render_sale(sale: Sale, language: str = "en") -> bytes:
        sale = Sale.objects.select_related("store", "customer", "created_by").prefetch_related("items", "ledger_entries").get(pk=sale.pk)
        local_created = timezone.localtime(sale.created_at, ZoneInfo(sale.store.timezone))
        total_balance = TransactionReceiptService._balance_after(sale, sale.outstanding if sale.customer else ZERO)
        return TransactionReceiptService._render(
            kind="sale",
            store=sale.store,
            number=TransactionReceiptService.sale_number(sale),
            counterparty_name=sale.customer.name if sale.customer else TransactionReceiptService._label("Walk-in customer", language),
            counterparty_phone=sale.customer.phone if sale.customer else "",
            created_at=f"{local_created:%Y-%m-%d %H:%M}",
            cashier=sale.created_by.get_username(),
            status=sale.status,
            items=[{
                "product_name": item.product_name,
                "factory_name": item.factory_name,
                "quantity": item.quantity,
                "pieces_per_unit": item.pieces_per_unit,
                "unit_rate": item.unit_price,
                "line_total": item.line_total,
            } for item in sale.items.all()],
            total=sale.total,
            amount_paid=sale.amount_paid,
            outstanding=sale.outstanding,
            total_balance=total_balance,
            note=sale.note,
            language=language,
        )

    @staticmethod
    def render_purchase(purchase: Purchase, language: str = "en") -> bytes:
        purchase = Purchase.objects.select_related("store", "supplier", "created_by").prefetch_related("items", "ledger_entries").get(pk=purchase.pk)
        total_balance = TransactionReceiptService._balance_after(purchase, -purchase.outstanding)
        return TransactionReceiptService._render(
            kind="purchase",
            store=purchase.store,
            number=TransactionReceiptService.purchase_number(purchase),
            counterparty_name=purchase.supplier.name,
            counterparty_phone=purchase.supplier.phone,
            created_at=str(purchase.purchase_date),
            cashier=purchase.created_by.get_username(),
            status=purchase.status,
            items=[{
                "product_name": item.product_name,
                "factory_name": item.factory_name,
                "quantity": item.quantity,
                "pieces_per_unit": item.pieces_per_unit,
                "unit_rate": item.unit_cost,
                "line_total": item.line_total,
            } for item in purchase.items.all()],
            total=purchase.total,
            amount_paid=purchase.amount_paid,
            outstanding=purchase.outstanding,
            total_balance=total_balance,
            reference=purchase.reference,
            note=purchase.note,
            language=language,
        )


class PaymentService:
    @staticmethod
    @transaction.atomic
    def create(*, store, user, party_id, amount, payment_date, method, idempotency_key, direction=Payment.Direction.RECEIVED, note=""):
        existing = Payment.objects.filter(store=store, idempotency_key=idempotency_key).first()
        if existing:
            return existing, False
        party = Party.objects.select_for_update().filter(store=store, is_active=True, id=party_id).first()
        if not party:
            raise ValidationError({"party_id": "Customer was not found in this store."})
        amount = money(amount)
        if amount <= 0:
            raise ValidationError({"amount": "Payment must be greater than zero."})
        if direction not in Payment.Direction.values:
            raise ValidationError({"direction": "Use received or sent."})
        if direction == Payment.Direction.SENT:
            amount_owed = abs(min(party.current_balance, ZERO))
            if amount_owed <= 0:
                raise ValidationError({"direction": "This party does not currently have an I Owe balance."})
            if amount > amount_owed:
                raise ValidationError({"amount": f"Payment sent cannot exceed the I Owe balance of {amount_owed}."})
        payment = Payment.objects.create(
            store=store,
            party=party,
            amount=amount,
            direction=direction,
            payment_date=payment_date,
            method=method,
            note=note,
            idempotency_key=idempotency_key,
            created_by=user,
        )
        delta = -amount if direction == Payment.Direction.RECEIVED else amount
        label = "Payment received" if direction == Payment.Direction.RECEIVED else "Payment sent"
        LedgerService.append(
            store=store,
            party=party,
            transaction_type=FinancialTransaction.TransactionType.PAYMENT,
            delta=delta,
            description=label,
            user=user,
            payment=payment,
            note=note,
        )
        AuditEvent.objects.create(
            store=store,
            user=user,
            action=f"payment.{direction}",
            entity_type="payment",
            entity_id=str(payment.id),
            details={"amount": str(amount), "party_id": party.id, "direction": direction},
        )
        return payment, True

    @staticmethod
    @transaction.atomic
    def adjustment(*, store, user, party_id, amount, direction, note):
        party = Party.objects.select_for_update().filter(store=store, is_active=True, id=party_id).first()
        if not party:
            raise ValidationError({"party_id": "Customer was not found in this store."})
        amount = money(amount)
        if amount <= 0:
            raise ValidationError({"amount": "Amount must be greater than zero."})
        if direction not in {"owes_me", "i_owe"}:
            raise ValidationError({"direction": "Use owes_me or i_owe."})
        delta = amount if direction == "owes_me" else -amount
        entry = LedgerService.append(
            store=store,
            party=party,
            transaction_type=FinancialTransaction.TransactionType.CREDIT,
            delta=delta,
            description="Credit or loan adjustment",
            user=user,
            note=note,
        )
        AuditEvent.objects.create(
            store=store,
            user=user,
            action="party.credit_adjusted",
            entity_type="party",
            entity_id=str(party.id),
            details={"direction": direction, "amount": str(amount), "ledger_entry_id": entry.id},
        )
        return entry


class ExpenseService:
    @staticmethod
    @transaction.atomic
    def create(*, store, user, category, amount, expense_date, payment_method, description, reference="", notes=""):
        if category.store_id != store.id or not category.is_active:
            raise ValidationError({"category": "Expense category was not found in this store."})
        amount = money(amount)
        if amount <= 0:
            raise ValidationError({"amount": "Expense amount must be greater than zero."})
        expense = Expense.objects.create(
            store=store,
            category=category,
            category_name=category.name,
            amount=amount,
            expense_date=expense_date,
            payment_method=payment_method,
            description=description,
            reference=reference,
            notes=notes,
            created_by=user,
        )
        AuditEvent.objects.create(
            store=store,
            user=user,
            action="expense.created",
            entity_type="expense",
            entity_id=str(expense.id),
            details={"amount": str(amount), "category": category.name},
        )
        return expense

    @staticmethod
    @transaction.atomic
    def reverse(*, expense, user, reason=""):
        locked = Expense.objects.select_for_update().get(pk=expense.pk, store=expense.store)
        if locked.is_reversed:
            return locked, False
        locked.is_reversed = True
        locked.reversed_at = timezone.now()
        locked.reversed_by = user
        locked.save(update_fields=["is_reversed", "reversed_at", "reversed_by", "updated_at"])
        AuditEvent.objects.create(
            store=locked.store,
            user=user,
            action="expense.reversed",
            entity_type="expense",
            entity_id=str(locked.id),
            details={"amount": str(locked.amount), "reason": reason},
        )
        return locked, True


class PurchaseService:
    @staticmethod
    @transaction.atomic
    def create(*, store, user, supplier_id, amount_paid, items, purchase_date, idempotency_key, reference="", note=""):
        existing = Purchase.objects.filter(store=store, idempotency_key=idempotency_key).prefetch_related("items").first()
        if existing:
            return existing, False
        if not items:
            raise ValidationError({"items": "At least one product is required."})
        supplier = Party.objects.select_for_update().filter(store=store, is_active=True, id=supplier_id).first()
        if not supplier:
            raise ValidationError({"supplier_id": "Factory was not found in this store."})
        if supplier.party_type != Party.PartyType.FACTORY:
            raise ValidationError({"supplier_id": "Choose a factory as the purchase supplier."})

        prepared_items = []
        seen_products = set()
        total = ZERO
        for item in items:
            product_id = item.get("product_id")
            if product_id in seen_products:
                raise ValidationError({"items": "A product may only appear once in a purchase."})
            seen_products.add(product_id)
            product = Product.objects.select_related("factory").filter(store=store, is_active=True, id=product_id).first()
            if not product:
                raise ValidationError({"items": f"Product {product_id} was not found."})
            if product.factory_id != supplier.id:
                raise ValidationError(
                    {"items": f"{product.name} belongs to {product.factory.name}, not {supplier.name}."}
                )
            item_quantity = quantity(item.get("quantity", 0))
            unit_cost = money(item.get("unit_cost", product.purchase_price))
            if item_quantity <= 0:
                raise ValidationError({"items": "Every quantity must be greater than zero."})
            if unit_cost < 0:
                raise ValidationError({"items": "Unit cost cannot be negative."})
            pieces_per_unit = product.pieces_per_unit
            line_total = money(item_quantity * pieces_per_unit * unit_cost)
            total += line_total
            prepared_items.append((product, item_quantity, pieces_per_unit, unit_cost, line_total))

        total = money(total)
        amount_paid = money(amount_paid)
        if amount_paid < 0 or amount_paid > total:
            raise ValidationError({"amount_paid": "Amount paid must be between zero and the purchase total."})
        purchase = Purchase.objects.create(
            store=store,
            supplier=supplier,
            idempotency_key=idempotency_key,
            purchase_date=purchase_date,
            subtotal=total,
            total=total,
            amount_paid=amount_paid,
            outstanding=money(total - amount_paid),
            reference=reference,
            note=note,
            created_by=user,
        )

        for product, item_quantity, pieces_per_unit, unit_cost, line_total in prepared_items:
            PurchaseItem.objects.create(
                purchase=purchase,
                product=product,
                product_name=product.name,
                factory_name=product.factory.name,
                quantity=item_quantity,
                pieces_per_unit=pieces_per_unit,
                unit_cost=unit_cost,
                line_total=line_total,
            )
            InventoryService.adjust(
                store=store,
                product=product,
                delta=item_quantity,
                movement_type=StockMovement.MovementType.RECEIVED,
                user=user,
                note=f"Purchase {purchase.id}",
                reference_type="purchase",
                reference_id=purchase.id,
            )
            if product.purchase_price != unit_cost:
                product.purchase_price = unit_cost
                product.save(update_fields=["purchase_price", "updated_at"])

        LedgerService.append(
            store=store,
            party=supplier,
            transaction_type=FinancialTransaction.TransactionType.PURCHASE,
            delta=-total,
            description=f"Purchase {str(purchase.id)[:8]}",
            user=user,
            purchase=purchase,
            note=note,
        )
        if amount_paid > 0:
            payment = Payment.objects.create(
                store=store,
                party=supplier,
                amount=amount_paid,
                direction=Payment.Direction.SENT,
                payment_date=purchase_date,
                method=Payment.Method.CASH,
                note=f"Payment recorded with purchase {purchase.id}",
                purchase=purchase,
                idempotency_key=uuid.uuid5(uuid.NAMESPACE_URL, f"purchase-payment:{store.id}:{purchase.id}"),
                created_by=user,
            )
            LedgerService.append(
                store=store,
                party=supplier,
                transaction_type=FinancialTransaction.TransactionType.PAYMENT,
                delta=amount_paid,
                description=f"Payment with purchase {str(purchase.id)[:8]}",
                user=user,
                payment=payment,
                purchase=purchase,
            )
        AuditEvent.objects.create(
            store=store,
            user=user,
            action="purchase.created",
            entity_type="purchase",
            entity_id=str(purchase.id),
            details={"total": str(total), "amount_paid": str(amount_paid), "supplier_id": supplier.id},
        )
        return purchase, True


def local_period_datetimes(store, period_start, period_end):
    tz = ZoneInfo(store.timezone)
    start_local = datetime.combine(period_start, time.min, tzinfo=tz)
    end_local = datetime.combine(period_end + timedelta(days=1), time.min, tzinfo=tz)
    return start_local.astimezone(ZoneInfo("UTC")), end_local.astimezone(ZoneInfo("UTC"))


class ReportService:
    @staticmethod
    def calculate(store, period_start, period_end):
        start_at, end_at = local_period_datetimes(store, period_start, period_end)
        sales = Sale.objects.filter(store=store, status=Sale.Status.COMPLETED, created_at__gte=start_at, created_at__lt=end_at)
        all_payments = Payment.objects.filter(
            store=store,
            direction=Payment.Direction.RECEIVED,
            created_at__gte=start_at,
            created_at__lt=end_at,
        )
        later_payments = all_payments.filter(sale__isnull=True, purchase__isnull=True)
        sent_payments = Payment.objects.filter(
            store=store,
            direction=Payment.Direction.SENT,
            created_at__gte=start_at,
            created_at__lt=end_at,
        )
        purchases = Purchase.objects.filter(
            store=store,
            status=Purchase.Status.COMPLETED,
            purchase_date__gte=period_start,
            purchase_date__lte=period_end,
        )
        expenses = Expense.objects.filter(
            store=store,
            is_reversed=False,
            expense_date__gte=period_start,
            expense_date__lte=period_end,
        )
        movements = StockMovement.objects.filter(store=store, created_at__gte=start_at, created_at__lt=end_at)

        sale_totals = sales.aggregate(
            total_sales=Coalesce(Sum("total"), ZERO),
            collected_at_sale=Coalesce(Sum("amount_paid"), ZERO),
            transaction_count=Count("id"),
        )
        credit_generated = sales.filter(outstanding__gt=0).aggregate(value=Coalesce(Sum("outstanding"), ZERO))["value"]
        profit_totals = SaleItem.objects.filter(sale__in=sales).aggregate(
            cost=Coalesce(Sum("line_cost"), ZERO),
            gross_profit=Coalesce(Sum("gross_profit"), ZERO),
        )
        later_collected = later_payments.aggregate(value=Coalesce(Sum("amount"), ZERO))["value"]
        purchase_totals = purchases.aggregate(
            total=Coalesce(Sum("total"), ZERO),
            paid=Coalesce(Sum("amount_paid"), ZERO),
            outstanding=Coalesce(Sum("outstanding"), ZERO),
            count=Count("id"),
        )
        expense_total = expenses.aggregate(value=Coalesce(Sum("amount"), ZERO))["value"]
        expense_categories = list(
            expenses.values("category_id", "category_name")
            .annotate(amount=Sum("amount"), count=Count("id"))
            .order_by("-amount")
        )
        receivables = Party.objects.filter(store=store, is_active=True, current_balance__gt=0).aggregate(value=Coalesce(Sum("current_balance"), ZERO))["value"]
        payables = Party.objects.filter(store=store, is_active=True, current_balance__lt=0).aggregate(value=Coalesce(Sum("current_balance"), ZERO))["value"]
        low_stock = list(
            Product.objects.filter(store=store, is_active=True, inventory__quantity__lte=F("minimum_stock_threshold"))
            .values("id", "name", factory_name=F("factory__name"), quantity=F("inventory__quantity"))
            .order_by("factory__name", "name")
        )
        products_sold = list(
            SaleItem.objects.filter(sale__in=sales)
            .values("product_id", "product_name", "factory_name")
            .annotate(
                quantity=Sum("quantity"),
                amount=Sum("line_total"),
                cost=Sum("line_cost"),
                gross_profit=Sum("gross_profit"),
            )
            .order_by("-quantity")
        )
        sold_product_ids = {item["product_id"] for item in products_sold}
        zero_sale_products = [
            {
                "product_id": item["id"],
                "product_name": item["name"],
                "factory_name": item["factory__name"],
                "quantity": Decimal("0.000"),
                "amount": ZERO,
                "cost": ZERO,
                "gross_profit": ZERO,
            }
            for item in Product.objects.filter(store=store, is_active=True)
            .exclude(id__in=sold_product_ids)
            .values("id", "name", "factory__name")
        ]
        lowest_selling_products = sorted([*zero_sale_products, *products_sold], key=lambda item: (item["quantity"], item["product_name"]))
        total_products_sold = sum((item["quantity"] for item in products_sold), Decimal("0.000"))
        stock_received_products = list(
            movements.filter(movement_type=StockMovement.MovementType.RECEIVED, quantity_delta__gt=0)
            .values("product_id", "product__name", "product__factory__name")
            .annotate(quantity=Sum("quantity_delta"))
            .order_by("-quantity")
        )
        raw_daily_rows = list(
            sales.annotate(day=TruncDate("created_at", tzinfo=ZoneInfo(store.timezone)))
            .values("day")
            .annotate(total=Sum("total"), collected=Sum("amount_paid"), transactions=Count("id"))
            .order_by("day")
        )
        days_in_period = max((period_end - period_start).days + 1, 1)
        raw_later_collections = {
            row["day"]: row["amount"]
            for row in later_payments.annotate(day=TruncDate("created_at", tzinfo=ZoneInfo(store.timezone)))
            .values("day")
            .annotate(amount=Sum("amount"))
        }
        raw_daily_by_date = {row["day"]: row for row in raw_daily_rows}
        daily_rows = []
        for day_offset in range(days_in_period):
            current_day = period_start + timedelta(days=day_offset)
            row = raw_daily_by_date.get(current_day, {})
            daily_rows.append(
                {
                    "day": current_day,
                    "total": row.get("total", ZERO),
                    "collected": row.get("collected", ZERO) + raw_later_collections.get(current_day, ZERO),
                    "transactions": row.get("transactions", 0),
                }
            )
        best_day = max(daily_rows, key=lambda row: row["total"], default=None)
        lowest_day = min(daily_rows, key=lambda row: row["total"], default=None)

        low_stock_event_map = {}
        for movement in movements.select_related("product", "product__factory").order_by("created_at", "id"):
            previous_balance = movement.balance_after - movement.quantity_delta
            reached_threshold = (
                movement.balance_after <= movement.product.minimum_stock_threshold
                and (
                    previous_balance > movement.product.minimum_stock_threshold
                    or movement.movement_type == StockMovement.MovementType.OPENING
                )
            )
            if not reached_threshold:
                continue
            item = low_stock_event_map.setdefault(
                movement.product_id,
                {
                    "product_id": movement.product_id,
                    "name": movement.product.name,
                    "factory_name": movement.product.factory.name,
                    "occurrences": 0,
                    "latest_quantity": str(movement.balance_after),
                },
            )
            item["occurrences"] += 1
            item["latest_quantity"] = str(movement.balance_after)
        low_stock_reached = sorted(low_stock_event_map.values(), key=lambda item: (-item["occurrences"], item["name"]))

        current_inventory = []
        for product in Product.objects.select_related("inventory", "factory").filter(store=store, is_active=True).order_by("factory__name", "name"):
            current_quantity = getattr(getattr(product, "inventory", None), "quantity", Decimal("0.000"))
            current_inventory.append(
                {
                    "product_id": product.id,
                    "name": product.name,
                    "factory_name": product.factory.name,
                    "unit": product.unit,
                    "quantity": str(current_quantity),
                    "purchase_price": str(money(product.purchase_price)),
                    "selling_price": str(money(product.selling_price)),
                    "is_low_stock": current_quantity <= product.minimum_stock_threshold,
                }
            )

        def party_summary(party_type):
            party_sales = sales.filter(customer__party_type=party_type)
            party_payments = all_payments.filter(party__party_type=party_type)
            active_parties = Party.objects.filter(store=store, is_active=True, party_type=party_type)
            totals = party_sales.aggregate(sales=Coalesce(Sum("total"), ZERO), transactions=Count("id"))
            received = party_payments.aggregate(value=Coalesce(Sum("amount"), ZERO))["value"]
            outstanding = active_parties.filter(current_balance__gt=0).aggregate(value=Coalesce(Sum("current_balance"), ZERO))["value"]
            credits = active_parties.filter(current_balance__lt=0).aggregate(value=Coalesce(Sum("current_balance"), ZERO))["value"]
            new_outstanding = party_sales.filter(outstanding__gt=0).aggregate(value=Coalesce(Sum("outstanding"), ZERO))["value"]
            top_sales = list(
                party_sales.values("customer_id", "customer__name")
                .annotate(amount=Sum("total"))
                .order_by("-amount")[:10]
            )
            top_outstanding = list(
                active_parties.filter(current_balance__gt=0)
                .values("id", "name", "current_balance")
                .order_by("-current_balance")[:10]
            )
            transaction_rows = []
            ledger_entries = (
                FinancialTransaction.objects.select_related("party")
                .filter(
                    store=store,
                    party__party_type=party_type,
                    created_at__gte=start_at,
                    created_at__lt=end_at,
                )
                .order_by("-created_at")[:250]
            )
            for entry in ledger_entries:
                transaction_rows.append(
                    {
                        "date_time": timezone.localtime(entry.created_at, ZoneInfo(store.timezone)).isoformat(),
                        "party_id": entry.party_id,
                        "party_name": entry.party.name,
                        "transaction_type": entry.transaction_type,
                        "description": entry.description,
                        "amount": str(abs(money(entry.delta))),
                        "credit_debit": "debit" if entry.delta > 0 else "credit",
                        "running_balance": str(money(entry.running_balance)),
                        "note": entry.note,
                    }
                )
            return {
                "sales": str(money(totals["sales"])),
                "transactions": totals["transactions"],
                "payments_received": str(money(received)),
                "new_outstanding": str(money(new_outstanding)),
                "outstanding": str(money(outstanding)),
                "customers_owing_count": active_parties.filter(current_balance__gt=0).count(),
                "store_owes": str(abs(money(credits))),
                "store_payables_count": active_parties.filter(current_balance__lt=0).count(),
                "top_by_sales": [
                    {"id": row["customer_id"], "name": row["customer__name"], "amount": str(money(row["amount"]))}
                    for row in top_sales
                ],
                "top_by_outstanding": [
                    {"id": row["id"], "name": row["name"], "amount": str(money(row["current_balance"]))}
                    for row in top_outstanding
                ],
                "transaction_history": transaction_rows,
            }

        stock_received = movements.filter(
            movement_type=StockMovement.MovementType.RECEIVED,
            quantity_delta__gt=0,
        ).aggregate(value=Coalesce(Sum("quantity_delta"), Decimal("0.000")))["value"]
        return {
            "period": {"start": str(period_start), "end": str(period_end), "timezone": store.timezone},
            "sales": {
                "total_sales": str(money(sale_totals["total_sales"])),
                "transaction_count": sale_totals["transaction_count"],
                "collected_at_sale": str(money(sale_totals["collected_at_sale"])),
                "debt_payments_received": str(money(later_collected)),
                "total_collected": str(money(sale_totals["collected_at_sale"] + later_collected)),
                "credit_generated": str(money(credit_generated)),
                "outstanding_generated": str(money(credit_generated)),
                "cost_of_goods_sold": str(money(profit_totals["cost"])),
                "gross_profit": str(money(profit_totals["gross_profit"])),
                "average_daily_sales": str(money(sale_totals["total_sales"] / days_in_period)),
                "best_sales_day": (
                    {"date": str(best_day["day"]), "amount": str(money(best_day["total"]))} if best_day else None
                ),
                "lowest_sales_day": (
                    {"date": str(lowest_day["day"]), "amount": str(money(lowest_day["total"]))} if lowest_day else None
                ),
                "daily_breakdown": [
                    {
                        "date": str(row["day"]),
                        "sales": str(money(row["total"])),
                        "collected": str(money(row["collected"])),
                        "transactions": row["transactions"],
                    }
                    for row in daily_rows
                ],
            },
            "financial_position": {
                "customer_receivables": str(money(receivables)),
                "customers_owing_count": Party.objects.filter(store=store, is_active=True, current_balance__gt=0).count(),
                "store_payables": str(abs(money(payables))),
                "store_payables_count": Party.objects.filter(store=store, is_active=True, current_balance__lt=0).count(),
                "payments_sent": str(money(sent_payments.aggregate(value=Coalesce(Sum("amount"), ZERO))["value"])),
            },
            "expenses": {
                "total": str(money(expense_total)),
                "transaction_count": expenses.count(),
                "by_category": [
                    {
                        "category_id": item["category_id"],
                        "category_name": item["category_name"],
                        "amount": str(money(item["amount"])),
                        "count": item["count"],
                    }
                    for item in expense_categories
                ],
            },
            "purchases": {
                "total": str(money(purchase_totals["total"])),
                "paid": str(money(purchase_totals["paid"])),
                "outstanding": str(money(purchase_totals["outstanding"])),
                "transaction_count": purchase_totals["count"],
            },
            "inventory": {
                "stock_received": str(quantity(stock_received)),
                "total_products_sold": str(quantity(total_products_sold)),
                "stock_received_products": [
                    {
                        "product_id": item["product_id"],
                        "name": item["product__name"],
                        "factory_name": item["product__factory__name"],
                        "quantity": str(item["quantity"]),
                    }
                    for item in stock_received_products
                ],
                "low_stock_occurrences": sum(item["occurrences"] for item in low_stock_reached),
                "products_reached_low_stock": low_stock_reached,
                "frequently_reaching_low_stock": low_stock_reached,
                "low_stock_count": len(low_stock),
                "low_stock": [{**item, "quantity": str(item["quantity"])} for item in low_stock],
                "products_sold": [
                    {
                        **item,
                        "quantity": str(item["quantity"]),
                        "amount": str(money(item["amount"])),
                        "cost": str(money(item["cost"])),
                        "gross_profit": str(money(item["gross_profit"])),
                    }
                    for item in products_sold
                ],
                "best_selling_products": [
                    {
                        **item,
                        "quantity": str(item["quantity"]),
                        "amount": str(money(item["amount"])),
                        "cost": str(money(item["cost"])),
                        "gross_profit": str(money(item["gross_profit"])),
                    }
                    for item in products_sold[:10]
                ],
                "lowest_selling_products": [
                    {
                        **item,
                        "quantity": str(item["quantity"]),
                        "amount": str(money(item["amount"])),
                        "cost": str(money(item["cost"])),
                        "gross_profit": str(money(item["gross_profit"])),
                    }
                    for item in lowest_selling_products[:10]
                ],
                "current_inventory": current_inventory,
            },
            "traders": party_summary(Party.PartyType.TRADER),
            "agents": party_summary(Party.PartyType.AGENT),
        }

    @staticmethod
    @transaction.atomic
    def generate(*, store, report_type, period_start, period_end, user=None):
        previous = Report.objects.filter(
            store=store,
            report_type=report_type,
            period_start=period_start,
            period_end=period_end,
        ).order_by("-version").first()
        version = (previous.version + 1) if previous else 1
        report = Report.objects.create(
            store=store,
            report_type=report_type,
            period_start=period_start,
            period_end=period_end,
            version=version,
            generated_by=user,
        )
        try:
            report.summary = ReportService.calculate(store, period_start, period_end)
            base_name = f"{store.id}-{report_type}-{period_start}-{period_end}-v{version}"
            report.pdf_file.save(f"{base_name}.pdf", ContentFile(ReportService.render_pdf(store, report.summary)), save=False)
            report.excel_file.save(f"{base_name}.xlsx", ContentFile(ReportService.render_excel(store, report.summary)), save=False)
            report.status = Report.Status.READY
            report.save()
            Notification.objects.create(
                store=store,
                notification_type=Notification.NotificationType.REPORT,
                title=f"{report.get_report_type_display()} report ready",
                message=f"Report for {period_start} to {period_end} is ready.",
                event_key=f"report:{report.id}",
                entity_type="report",
                entity_id=str(report.id),
            )
        except Exception as exc:
            report.status = Report.Status.FAILED
            report.error_message = str(exc)
            report.save(update_fields=["status", "error_message", "updated_at"])
            raise
        return report

    @staticmethod
    def render_pdf(store, summary):
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=A4)
        _, height = A4
        y = height - 48

        def page_header():
            nonlocal y
            pdf.setFont("Helvetica-Bold", 16)
            pdf.drawString(42, y, f"{store.name} - Business Report")
            y -= 23
            pdf.setFont("Helvetica", 9)
            pdf.drawString(42, y, f"Period: {summary['period']['start']} to {summary['period']['end']} ({summary['period']['timezone']})")
            y -= 24

        def ensure_space(required=30):
            nonlocal y
            if y - required < 48:
                pdf.showPage()
                y = height - 48
                page_header()

        def heading(title):
            nonlocal y
            ensure_space(34)
            pdf.setFont("Helvetica-Bold", 12)
            pdf.drawString(42, y, title)
            y -= 18

        def line(text, indent=0, font="Helvetica", size=9):
            nonlocal y
            ensure_space(16)
            pdf.setFont(font, size)
            clean = str(text).replace("\n", " ")[:118]
            pdf.drawString(48 + indent, y, clean)
            y -= 14

        def scalar_section(title, values):
            nonlocal y
            heading(title)
            for key, value in values.items():
                if isinstance(value, dict):
                    nested = ", ".join(f"{nested_key.replace('_', ' ')}: {nested_value}" for nested_key, nested_value in value.items())
                    line(f"{key.replace('_', ' ').title()}: {nested}", indent=8)
                elif not isinstance(value, list):
                    line(f"{key.replace('_', ' ').title()}: {value}", indent=8)
            y -= 5

        page_header()
        scalar_section("Sales Summary", summary["sales"])
        scalar_section("Financial Position", summary["financial_position"])
        scalar_section("Expenses", summary["expenses"])
        scalar_section("Purchases", summary["purchases"])
        scalar_section("Inventory Summary", summary["inventory"])

        heading("Daily Sales Breakdown")
        line("Date | Sales | Collected | Transactions", font="Helvetica-Bold")
        for item in summary["sales"]["daily_breakdown"]:
            line(f"{item['date']} | {item['sales']} | {item['collected']} | {item['transactions']}")

        def product_table(title, rows, quantity_key="quantity"):
            heading(title)
            if not rows:
                line("No records for this period.")
                return
            for item in rows:
                name = item.get("product_name") or item.get("name") or "Product"
                factory = item.get("factory_name", "")
                amount = f" | {item['amount']} {store.currency}" if "amount" in item else ""
                profit = f" | profit {item['gross_profit']} {store.currency}" if "gross_profit" in item else ""
                line(f"{name} (Factory: {factory}) | {item.get(quantity_key, '')}{amount}{profit}")

        inventory = summary["inventory"]
        product_table("Products Sold", inventory["products_sold"])
        product_table("Stock Received", inventory["stock_received_products"])
        product_table("Current Low Stock", inventory["low_stock"])
        product_table("Current Inventory", inventory["current_inventory"])

        for party_type, title in (("traders", "Trader Summary"), ("agents", "Agent Summary")):
            values = summary[party_type]
            scalar_section(title, values)
            heading(f"{title} - Top by Sales")
            for item in values["top_by_sales"]:
                line(f"{item['name']} | {item['amount']} {store.currency}")
            heading(f"{title} - Top Outstanding")
            for item in values["top_by_outstanding"]:
                line(f"{item['name']} | {item['amount']} {store.currency}")
            heading(f"{title} - Transaction History")
            if not values["transaction_history"]:
                line("No transactions for this period.")
            for item in values["transaction_history"]:
                line(
                    f"{item['date_time'][:16]} | {item['party_name']} | {item['transaction_type']} | "
                    f"{item['amount']} {store.currency} | balance {item['running_balance']}"
                )

        pdf.save()
        return buffer.getvalue()

    @staticmethod
    def render_excel(store, summary):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Summary"
        sheet.append([store.name, "Business Report"])
        sheet.append(["Period", f"{summary['period']['start']} to {summary['period']['end']}"])
        sheet.append([])
        for section_name in ("sales", "financial_position", "expenses", "purchases"):
            sheet.append([section_name.replace("_", " ").title()])
            for key, value in summary[section_name].items():
                if isinstance(value, dict):
                    sheet.append([key.replace("_", " ").title(), ", ".join(f"{nested_key}: {nested_value}" for nested_key, nested_value in value.items())])
                    continue
                if isinstance(value, list):
                    continue
                sheet.append([key.replace("_", " ").title(), value])
            sheet.append([])
        inventory = summary["inventory"]

        def product_sheet(title, rows):
            target = workbook.create_sheet(title)
            target.append(["Product", "Factory", "Unit", "Quantity", "Amount", "Cost", "Gross Profit", "Purchase Price", "Selling Price", "Low Stock"])
            for item in rows:
                target.append(
                    [
                        item.get("product_name") or item.get("name"),
                        item.get("factory_name", ""),
                        item.get("unit", ""),
                        item.get("quantity", ""),
                        item.get("amount", ""),
                        item.get("cost", ""),
                        item.get("gross_profit", ""),
                        item.get("purchase_price", ""),
                        item.get("selling_price", ""),
                        item.get("is_low_stock", ""),
                    ]
                )

        product_sheet("Products Sold", inventory["products_sold"])
        product_sheet("Stock Received", inventory["stock_received_products"])
        product_sheet("Low Stock", inventory["low_stock"])
        product_sheet("Current Inventory", inventory["current_inventory"])
        reached = workbook.create_sheet("Low Stock Events")
        reached.append(["Product", "Factory", "Occurrences", "Latest Quantity"])
        for item in inventory["products_reached_low_stock"]:
            reached.append([item["name"], item["factory_name"], item["occurrences"], item["latest_quantity"]])
        daily = workbook.create_sheet("Daily Sales")
        daily.append(["Date", "Sales", "Collected", "Transactions"])
        for item in summary["sales"]["daily_breakdown"]:
            daily.append([item["date"], item["sales"], item["collected"], item["transactions"]])
        customers = workbook.create_sheet("Customer Summary")
        customers.append(["Type", "Sales", "Transactions", "Payments", "New Outstanding", "Outstanding", "Owing Count", "Store Owes", "Payables Count"])
        for label in ("traders", "agents"):
            values = summary[label]
            customers.append(
                [label.title(), values["sales"], values["transactions"], values["payments_received"], values["new_outstanding"], values["outstanding"], values["customers_owing_count"], values["store_owes"], values["store_payables_count"]]
            )
        top = workbook.create_sheet("Top Customers")
        top.append(["Type", "Ranking", "Customer", "Amount"])
        history = workbook.create_sheet("Customer Transactions")
        history.append(["Type", "Date/Time", "Customer", "Transaction Type", "Description", "Amount", "Credit/Debit", "Running Balance", "Note"])
        for label in ("traders", "agents"):
            values = summary[label]
            for item in values["top_by_sales"]:
                top.append([label.title(), "Sales", item["name"], item["amount"]])
            for item in values["top_by_outstanding"]:
                top.append([label.title(), "Outstanding", item["name"], item["amount"]])
            for item in values["transaction_history"]:
                history.append(
                    [label.title(), item["date_time"], item["party_name"], item["transaction_type"], item["description"], item["amount"], item["credit_debit"], item["running_balance"], item["note"]]
                )
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            for column in worksheet.columns:
                letter = column[0].column_letter
                worksheet.column_dimensions[letter].width = min(max((len(str(cell.value or "")) for cell in column), default=8) + 2, 45)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()


class SMSService:
    @staticmethod
    def _send(phone, message, provider):
        if provider == "console" or settings.SMS_PROVIDER == "console":
            print(f"SMS to {phone}: {message}")
            return f"console-{uuid.uuid4()}"
        if not settings.SMS_API_URL or not settings.SMS_API_TOKEN:
            raise RuntimeError("SMS_API_URL and SMS_API_TOKEN must be configured.")
        response = requests.post(
            settings.SMS_API_URL,
            json={"to": phone, "message": message, "sender_id": settings.SMS_SENDER_ID},
            headers={"Authorization": f"Bearer {settings.SMS_API_TOKEN}"},
            timeout=15,
        )
        response.raise_for_status()
        payload = response.json() if response.content else {}
        return str(payload.get("id", ""))

    @staticmethod
    def send_debt_reminders(store, reminder_date):
        store_settings, _ = StoreSettings.objects.get_or_create(store=store)
        if not store_settings.sms_enabled:
            return {"sent": 0, "failed": 0, "disabled": True}
        result = {"sent": 0, "failed": 0, "disabled": False}
        parties = (
            Party.objects.filter(store=store, is_active=True, sms_enabled=True, current_balance__gt=0)
            .exclude(party_type=Party.PartyType.FACTORY)
            .exclude(phone="")
        )
        account_number = store_settings.sms_account_number or store.account_number
        for party in parties:
            message = (
                f"Dear {party.name}, your outstanding balance is {party.current_balance} {store.currency}. "
                f"Please make payment to Account No. {account_number}. Thank you."
            )
            try:
                log, created = SMSLog.objects.get_or_create(
                    store=store,
                    party=party,
                    reminder_date=reminder_date,
                    defaults={
                        "phone": party.phone,
                        "outstanding_amount": party.current_balance,
                        "message": message,
                        "provider": store_settings.sms_provider,
                    },
                )
            except IntegrityError:
                continue
            if not created or log.status == SMSLog.Status.SENT:
                continue
            try:
                log.provider_message_id = SMSService._send(party.phone, message, store_settings.sms_provider)
                log.status = SMSLog.Status.SENT
                log.sent_at = timezone.now()
                result["sent"] += 1
            except Exception as exc:
                log.status = SMSLog.Status.FAILED
                log.error_message = str(exc)
                result["failed"] += 1
            log.save(update_fields=["provider_message_id", "status", "sent_at", "error_message"])
        return result
