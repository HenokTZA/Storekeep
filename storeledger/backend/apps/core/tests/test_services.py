import uuid
import tempfile
from io import BytesIO
from datetime import date
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from apps.core.models import (
    Expense,
    ExpenseCategory,
    FinancialTransaction,
    InventoryBalance,
    Membership,
    Notification,
    Party,
    Product,
    Purchase,
    SMSLog,
    Store,
    StoreSettings,
)
from apps.core.services import ExpenseService, InventoryService, PaymentService, PurchaseService, ReportService, SMSService, SaleService
from apps.core.models import StockMovement


class StoreServiceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("owner", password="test-password")
        self.store = Store.objects.create(name="Test Store", account_number="123", currency="ETB")
        Membership.objects.create(user=self.user, store=self.store, role=Membership.Role.OWNER)
        self.store_settings = StoreSettings.objects.create(store=self.store)
        self.product = Product.objects.create(
            store=self.store,
            name="LED Driver",
            sku="LED-1",
            selling_price=Decimal("50.00"),
            minimum_stock_threshold=Decimal("5.000"),
        )
        InventoryService.adjust(
            store=self.store,
            product=self.product,
            delta="10",
            movement_type=StockMovement.MovementType.OPENING,
            user=self.user,
        )
        self.party = Party.objects.create(
            store=self.store,
            party_type=Party.PartyType.TRADER,
            name="Ahmed Trading",
            phone="+251911000000",
        )

    def create_sale(self, key=None, quantity="2", paid="40"):
        return SaleService.create(
            store=self.store,
            user=self.user,
            customer_id=self.party.id,
            amount_paid=paid,
            items=[{"product_id": self.product.id, "quantity": quantity}],
            idempotency_key=key or uuid.uuid4(),
        )

    def test_partial_sale_updates_stock_and_balance_atomically(self):
        sale, created = self.create_sale()
        self.assertTrue(created)
        self.assertEqual(sale.total, Decimal("100.00"))
        self.assertEqual(sale.outstanding, Decimal("60.00"))
        self.assertEqual(InventoryBalance.objects.get(product=self.product).quantity, Decimal("8.000"))
        self.party.refresh_from_db()
        self.assertEqual(self.party.current_balance, Decimal("60.00"))
        self.assertEqual(FinancialTransaction.objects.filter(party=self.party).count(), 2)

    def test_idempotent_sale_does_not_double_post(self):
        key = uuid.uuid4()
        first, first_created = self.create_sale(key=key)
        second, second_created = self.create_sale(key=key)
        self.assertTrue(first_created)
        self.assertFalse(second_created)
        self.assertEqual(first.id, second.id)
        self.assertEqual(InventoryBalance.objects.get(product=self.product).quantity, Decimal("8.000"))

    def test_agent_sale_always_uses_one_and_a_half_percent_discount(self):
        agent = Party.objects.create(
            store=self.store,
            party_type=Party.PartyType.AGENT,
            name="Agent Customer",
            phone="+251911000001",
        )
        sale, _ = SaleService.create(
            store=self.store,
            user=self.user,
            customer_id=agent.id,
            amount_paid="98.50",
            items=[{"product_id": self.product.id, "quantity": "2", "unit_price": "999.00"}],
            idempotency_key=uuid.uuid4(),
        )
        self.assertEqual(sale.items.get().unit_price, Decimal("49.25"))
        self.assertEqual(sale.total, Decimal("98.50"))

    def test_pack_sale_uses_units_times_pieces_times_piece_price(self):
        self.product.selling_price = Decimal("200.00")
        self.product.pieces_per_unit = 54
        self.product.save(update_fields=["selling_price", "pieces_per_unit", "updated_at"])
        sale, _ = self.create_sale(quantity="1", paid="10800")
        item = sale.items.get()
        self.assertEqual(sale.total, Decimal("10800.00"))
        self.assertEqual(item.quantity, Decimal("1.000"))
        self.assertEqual(item.pieces_per_unit, 54)
        self.assertEqual(item.unit_price, Decimal("200.00"))
        self.assertEqual(item.line_total, Decimal("10800.00"))
        self.assertEqual(InventoryBalance.objects.get(product=self.product).quantity, Decimal("9.000"))

        self.product.pieces_per_unit = 60
        self.product.save(update_fields=["pieces_per_unit", "updated_at"])
        item.refresh_from_db()
        self.assertEqual(item.pieces_per_unit, 54, "Completed sales must keep their original pack snapshot.")

    def test_negative_inventory_is_rejected_and_sale_rolls_back(self):
        with self.assertRaises(ValidationError):
            self.create_sale(quantity="50")
        self.assertEqual(InventoryBalance.objects.get(product=self.product).quantity, Decimal("10.000"))
        self.assertEqual(self.store.sales.count(), 0)

    def test_payment_can_create_customer_credit(self):
        self.create_sale(paid="0")
        PaymentService.create(
            store=self.store,
            user=self.user,
            party_id=self.party.id,
            amount="125",
            payment_date=date.today(),
            method="cash",
            idempotency_key=uuid.uuid4(),
        )
        self.party.refresh_from_db()
        self.assertEqual(self.party.current_balance, Decimal("-25.00"))
        self.assertEqual(self.party.balance_label, "I Owe")

    def test_payment_sent_settles_i_owe_balance_without_overpayment(self):
        PaymentService.adjustment(
            store=self.store,
            user=self.user,
            party_id=self.party.id,
            amount="75",
            direction="i_owe",
            note="Supplier credit",
        )
        payment, created = PaymentService.create(
            store=self.store,
            user=self.user,
            party_id=self.party.id,
            amount="50",
            payment_date=date.today(),
            method="bank",
            direction="sent",
            idempotency_key=uuid.uuid4(),
        )
        self.assertTrue(created)
        self.assertEqual(payment.direction, "sent")
        self.party.refresh_from_db()
        self.assertEqual(self.party.current_balance, Decimal("-25.00"))
        with self.assertRaises(ValidationError):
            PaymentService.create(
                store=self.store,
                user=self.user,
                party_id=self.party.id,
                amount="30",
                payment_date=date.today(),
                method="cash",
                direction="sent",
                idempotency_key=uuid.uuid4(),
            )

    def test_purchase_posts_stock_payable_and_payment_atomically(self):
        key = uuid.uuid4()
        purchase, created = PurchaseService.create(
            store=self.store,
            user=self.user,
            supplier_id=self.party.id,
            amount_paid="40",
            items=[{"product_id": self.product.id, "quantity": "2", "unit_cost": "30"}],
            purchase_date=date.today(),
            idempotency_key=key,
        )
        duplicate, duplicate_created = PurchaseService.create(
            store=self.store,
            user=self.user,
            supplier_id=self.party.id,
            amount_paid="40",
            items=[{"product_id": self.product.id, "quantity": "2", "unit_cost": "30"}],
            purchase_date=date.today(),
            idempotency_key=key,
        )
        self.assertTrue(created)
        self.assertFalse(duplicate_created)
        self.assertEqual(purchase.id, duplicate.id)
        self.assertEqual(Purchase.objects.count(), 1)
        self.assertEqual(InventoryBalance.objects.get(product=self.product).quantity, Decimal("12.000"))
        self.party.refresh_from_db()
        self.assertEqual(self.party.current_balance, Decimal("-20.00"))

    def test_pack_purchase_uses_piece_cost_and_receives_pack_units(self):
        self.product.pieces_per_unit = 50
        self.product.save(update_fields=["pieces_per_unit", "updated_at"])
        purchase, _ = PurchaseService.create(
            store=self.store,
            user=self.user,
            supplier_id=self.party.id,
            amount_paid="3000",
            items=[{"product_id": self.product.id, "quantity": "2", "unit_cost": "30"}],
            purchase_date=date.today(),
            idempotency_key=uuid.uuid4(),
        )
        item = purchase.items.get()
        self.assertEqual(purchase.total, Decimal("3000.00"))
        self.assertEqual(item.pieces_per_unit, 50)
        self.assertEqual(item.line_total, Decimal("3000.00"))
        self.assertEqual(InventoryBalance.objects.get(product=self.product).quantity, Decimal("12.000"))

    def test_expense_is_snapshot_based_and_reversible_without_deletion(self):
        category = ExpenseCategory.objects.create(store=self.store, name="Transport")
        expense = ExpenseService.create(
            store=self.store,
            user=self.user,
            category=category,
            amount="125.50",
            expense_date=date.today(),
            payment_method="cash",
            description="Delivery",
        )
        category.name = "Delivery & transport"
        category.save()
        expense.refresh_from_db()
        self.assertEqual(expense.category_name, "Transport")
        reversed_expense, created = ExpenseService.reverse(expense=expense, user=self.user, reason="Duplicate")
        self.assertTrue(created)
        self.assertTrue(reversed_expense.is_reversed)
        self.assertEqual(Expense.objects.count(), 1)

    def test_low_stock_notification_is_deduplicated(self):
        self.create_sale(quantity="5", paid="250")
        self.assertEqual(Notification.objects.filter(event_key=f"low-stock:{self.product.id}").count(), 1)
        InventoryService.adjust(
            store=self.store,
            product=self.product,
            delta="-1",
            movement_type=StockMovement.MovementType.ADJUSTMENT,
            user=self.user,
        )
        self.assertEqual(Notification.objects.filter(event_key=f"low-stock:{self.product.id}").count(), 1)

    @override_settings(SMS_PROVIDER="console")
    def test_sms_only_targets_positive_balances_once_per_day(self):
        self.create_sale(paid="0")
        self.store_settings.sms_enabled = True
        self.store_settings.save()
        result = SMSService.send_debt_reminders(self.store, date.today())
        second = SMSService.send_debt_reminders(self.store, date.today())
        self.assertEqual(result["sent"], 1)
        self.assertEqual(second["sent"], 0)
        self.assertEqual(SMSLog.objects.count(), 1)

    def test_report_uses_recorded_transactions(self):
        self.create_sale()
        unused = Product.objects.create(store=self.store, name="Unsold Product", sku="ZERO-1", selling_price=Decimal("20.00"))
        InventoryBalance.objects.create(product=unused, quantity=Decimal("7.000"))
        summary = ReportService.calculate(self.store, date.today(), date.today())
        self.assertEqual(summary["sales"]["total_sales"], "100.00")
        self.assertEqual(summary["sales"]["collected_at_sale"], "40.00")
        self.assertEqual(summary["sales"]["credit_generated"], "60.00")
        self.assertEqual(summary["inventory"]["total_products_sold"], "2.000")
        self.assertEqual(summary["inventory"]["products_sold"][0]["product_name"], "LED Driver")
        self.assertEqual(len(summary["inventory"]["current_inventory"]), 2)
        self.assertEqual(summary["inventory"]["lowest_selling_products"][0]["product_name"], "Unsold Product")
        self.assertEqual(summary["traders"]["customers_owing_count"], 1)
        self.assertEqual(len(summary["traders"]["transaction_history"]), 2)
        self.assertIn("expenses", summary)
        self.assertIn("purchases", summary)

    def test_report_generation_writes_pdf_and_excel(self):
        self.create_sale()
        with tempfile.TemporaryDirectory() as media_root, self.settings(MEDIA_ROOT=media_root):
            report = ReportService.generate(
                store=self.store,
                report_type="custom",
                period_start=date.today(),
                period_end=date.today(),
                user=self.user,
            )
            self.assertEqual(report.status, "ready")
            self.assertTrue(report.pdf_file.name.endswith(".pdf"))
            self.assertTrue(report.excel_file.name.endswith(".xlsx"))
            self.assertGreater(report.pdf_file.size, 500)
            self.assertGreater(report.excel_file.size, 500)
            workbook = load_workbook(BytesIO(report.excel_file.read()), read_only=True)
            self.assertTrue(
                {"Summary", "Daily Sales", "Products Sold", "Stock Received", "Low Stock", "Current Inventory", "Low Stock Events", "Customer Summary", "Top Customers", "Customer Transactions"}.issubset(workbook.sheetnames)
            )
